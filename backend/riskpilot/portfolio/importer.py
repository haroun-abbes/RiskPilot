"""Portfolio file import: parsing, column mapping, validation.

Workflow: preview (parse + suggest mapping) -> validate (mapping applied,
row-level errors) -> commit (only when no blocking error remains).

Field conventions (documented in the template):
  coupon_rate in PERCENT (3.25 means 3.25 %), clean_price in % of par,
  quantity in shares/units (negative = short), nominal in currency units,
  dates as YYYY-MM-DD or DD/MM/YYYY.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime

from ..quant.daycount import SUPPORTED_FREQUENCIES, normalize_daycount
from ..quant.instruments import (
    SUPPORTED_TYPES,
    UNSUPPORTED_TYPES_MESSAGES,
    EquityPosition,
    FixedBondPosition,
    FxCashPosition,
    rating_bucket,
)

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
ISO_CCY = {"EUR", "USD", "GBP", "CHF", "JPY", "SEK", "NOK", "DKK", "CAD", "AUD", "HKD", "SGD", "CNY", "PLN", "CZK", "HUF"}

FIELDS = {
    "id": "Unique position identifier",
    "instrument_type": "EQUITY | ETF | BOND_FIXED | FX_CASH",
    "ticker": "Equity/ETF ticker (matches EQ:<ticker> in market data)",
    "name": "Display name",
    "quantity": "Shares or units (negative = short)",
    "price": "Price per share in local currency",
    "nominal": "Bond face amount (negative = short)",
    "clean_price": "Bond clean price in % of par",
    "coupon_rate": "Annual coupon in percent",
    "frequency": "Coupons per year (1, 2, 4, 12)",
    "maturity": "Bond maturity date",
    "daycount": "ACT/ACT ICMA, 30E/360, 30/360, ACT/365F, ACT/360",
    "currency": "ISO 4217 code",
    "amount": "Cash amount in currency (FX_CASH)",
    "sector": "Sector",
    "issuer": "Issuer",
    "country": "Country (ISO 2 letters)",
    "rating": "Rating (S&P scale)",
    "is_sovereign": "TRUE for government bonds (no credit-spread factor)",
    "risk_proxy": "Ticker of a proxy price series when no own history",
}
REQUIRED = {
    "EQUITY": ["id", "ticker", "quantity", "price", "currency"],
    "ETF": ["id", "ticker", "quantity", "price", "currency"],
    "BOND_FIXED": ["id", "nominal", "clean_price", "coupon_rate", "frequency", "maturity", "currency"],
    "FX_CASH": ["id", "currency", "amount"],
}
SYNONYMS = {
    "id": ["id", "position_id", "positionid", "identifiant", "isin", "code"],
    "instrument_type": ["instrument_type", "type", "asset_type", "assettype", "instrument", "classe"],
    "ticker": ["ticker", "symbol", "symbole", "mnemo"],
    "name": ["name", "nom", "description", "libelle", "security_name"],
    "quantity": ["quantity", "qty", "quantite", "shares", "units"],
    "price": ["price", "prix", "last_price", "cours"],
    "nominal": ["nominal", "face", "face_value", "par_amount", "notional"],
    "clean_price": ["clean_price", "cleanprice", "prix_pied_de_coupon", "bond_price"],
    "coupon_rate": ["coupon_rate", "coupon", "taux_coupon", "cpn"],
    "frequency": ["frequency", "freq", "frequence", "coupon_frequency"],
    "maturity": ["maturity", "maturity_date", "echeance", "date_echeance"],
    "daycount": ["daycount", "day_count", "base", "convention", "basis"],
    "currency": ["currency", "ccy", "devise", "curr"],
    "amount": ["amount", "montant", "cash", "balance"],
    "sector": ["sector", "secteur", "industry"],
    "issuer": ["issuer", "emetteur", "issuer_name"],
    "country": ["country", "pays", "country_code"],
    "rating": ["rating", "notation", "credit_rating"],
    "is_sovereign": ["is_sovereign", "sovereign", "souverain", "govt"],
    "risk_proxy": ["risk_proxy", "proxy", "proxy_ticker"],
}


class ImportFileError(ValueError):
    pass


def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower().strip()
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_")


def parse_file(filename: str, raw: bytes) -> tuple[list[str], list[list]]:
    if len(raw) == 0:
        raise ImportFileError("The file is empty.")
    if len(raw) > MAX_BYTES:
        raise ImportFileError(f"File exceeds {MAX_BYTES // 1024 // 1024} MB.")
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext in ("csv", "txt"):
        rows = _parse_csv(raw)
    elif ext == "xlsx":
        rows = _parse_xlsx(raw)
    else:
        raise ImportFileError("Unsupported file type. Use .csv or .xlsx (macro-enabled and legacy .xls files are refused).")
    rows = [r for r in rows if any(str(c).strip() for c in r if c is not None)]
    if len(rows) < 2:
        raise ImportFileError("The file needs a header row and at least one position.")
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    if any(h == "" for h in header):
        raise ImportFileError("Empty column header found; every column needs a name.")
    if len(set(_norm(h) for h in header)) != len(header):
        raise ImportFileError("Duplicate column headers.")
    body = rows[1:]
    if len(body) > MAX_ROWS:
        raise ImportFileError(f"Too many rows ({len(body)} > {MAX_ROWS}).")
    width = len(header)
    body = [list(r[:width]) + [None] * (width - len(r)) for r in body]
    return header, body


def _parse_csv(raw: bytes) -> list[list]:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = raw.decode("cp1252")
        except UnicodeDecodeError:
            raise ImportFileError("Unrecognised text encoding (use UTF-8).")
    if "\x00" in text:
        raise ImportFileError("Binary content detected; not a CSV file.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        delim = dialect.delimiter
    except csv.Error:
        delim = ","
    return [row for row in csv.reader(io.StringIO(text), delimiter=delim)]


def _parse_xlsx(raw: bytes) -> list[list]:
    import zipfile

    from openpyxl import load_workbook

    if not zipfile.is_zipfile(io.BytesIO(raw)):
        raise ImportFileError("Not a valid .xlsx file.")
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:
        raise ImportFileError(f"Unreadable Excel workbook: {type(e).__name__}") from e
    ws = wb["Positions"] if "Positions" in wb.sheetnames else wb.worksheets[0]
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i > MAX_ROWS + 1:
            raise ImportFileError(f"Too many rows (> {MAX_ROWS}).")
        rows.append(list(r))
    wb.close()
    return rows


def suggest_mapping(header: list[str]) -> dict[str, str | None]:
    """Canonical field -> source column (or None)."""
    normed = {_norm(h): h for h in header}
    out = {}
    used = set()
    for field, syns in SYNONYMS.items():
        hit = next((normed[s] for s in syns if s in normed and normed[s] not in used), None)
        out[field] = hit
        if hit:
            used.add(hit)
    return out


# ---------- value parsing ----------
def _num(v):
    if v is None or (isinstance(v, str) and v.strip() == ""):
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    s = str(v).strip().replace(" ", "").replace(" ", "").replace("%", "")
    if "," in s and "." in s:
        s = s.replace(",", "") if s.rfind(".") > s.rfind(",") else s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    return float(s)


def _date(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"'{s}' is not a recognised date (YYYY-MM-DD or DD/MM/YYYY)")


def _bool(v):
    if v is None or str(v).strip() == "":
        return False
    s = str(v).strip().lower()
    if s in ("true", "1", "yes", "oui", "y", "vrai"):
        return True
    if s in ("false", "0", "no", "non", "n", "faux"):
        return False
    raise ValueError(f"'{v}' is not a boolean (TRUE/FALSE)")


def _str(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


@dataclass
class RowIssue:
    row: int  # 1-based data row (header excluded)
    field: str | None
    message: str
    severity: str  # error | warning

    def dict(self):
        return self.__dict__


def validate(header, body, mapping: dict[str, str | None], as_of: date):
    """Return (positions, issues, rows_preview). Blocking issues have severity 'error'."""
    col = {h: i for i, h in enumerate(header)}
    for f, src in mapping.items():
        if f not in FIELDS:
            raise ImportFileError(f"Unknown target field '{f}'")
        if src is not None and src not in col:
            raise ImportFileError(f"Mapped column '{src}' does not exist in the file")
    if not mapping.get("instrument_type"):
        raise ImportFileError("The 'instrument_type' field must be mapped.")
    if not mapping.get("id"):
        raise ImportFileError("The 'id' field must be mapped.")

    issues: list[RowIssue] = []
    positions = []
    seen: dict[str, int] = {}
    preview = []
    for r_i, row in enumerate(body, start=1):
        rec = {f: (row[col[src]] if src else None) for f, src in mapping.items()}
        status = "ok"

        def err(field, msg):
            nonlocal status
            issues.append(RowIssue(r_i, field, msg, "error"))
            status = "error"

        def warn(field, msg):
            nonlocal status
            issues.append(RowIssue(r_i, field, msg, "warning"))
            if status == "ok":
                status = "warning"

        itype = (_str(rec.get("instrument_type")) or "").upper()
        pid = _str(rec.get("id"))
        if not pid:
            err("id", "Missing position id")
        elif pid in seen:
            err("id", f"Duplicate id '{pid}' (also on row {seen[pid]})")
        else:
            seen[pid] = r_i
        if itype not in SUPPORTED_TYPES:
            msg = UNSUPPORTED_TYPES_MESSAGES.get(itype, f"Unknown instrument type '{itype or '(empty)'}'. Supported: {', '.join(SUPPORTED_TYPES)}")
            err("instrument_type", msg)
            preview.append({"row": r_i, "id": pid, "instrument_type": itype, "status": status})
            continue
        for f in REQUIRED[itype]:
            if mapping.get(f) is None:
                err(f, f"Field '{f}' is required for {itype} but no column is mapped")
            elif _str(rec.get(f)) is None:
                err(f, f"Missing value for required field '{f}'")
        ccy = (_str(rec.get("currency")) or "").upper()
        if ccy and ccy not in ISO_CCY:
            err("currency", f"Unknown or unsupported currency '{ccy}'")

        vals = {}
        for f in ("quantity", "price", "nominal", "clean_price", "coupon_rate", "frequency", "amount"):
            try:
                vals[f] = _num(rec.get(f))
            except (ValueError, TypeError):
                err(f, f"'{rec.get(f)}' is not a number")
                vals[f] = None
        try:
            mat = _date(rec.get("maturity"))
        except ValueError as e:
            err("maturity", str(e))
            mat = None
        try:
            sov = _bool(rec.get("is_sovereign"))
        except ValueError as e:
            err("is_sovereign", str(e))
            sov = False

        pos = None
        common = dict(name=_str(rec.get("name")))
        # every check runs (so all problems of a row are reported at once);
        # a position object is only built when the row has no blocking error
        if True:
            try:
                if itype in ("EQUITY", "ETF"):
                    if vals["quantity"] == 0:
                        err("quantity", "Quantity is zero")
                    if vals["price"] is not None and vals["price"] <= 0:
                        err("price", "Price must be strictly positive")
                    if status != "error":
                        pos = EquityPosition(
                            id=pid, ticker=_str(rec["ticker"]).upper(), quantity=vals["quantity"], price=vals["price"],
                            currency=ccy, instrument_type=itype, sector=_str(rec.get("sector")), issuer=_str(rec.get("issuer")),
                            country=_str(rec.get("country")), risk_proxy=(_str(rec.get("risk_proxy")) or "").upper() or None, **common,
                        )
                elif itype == "BOND_FIXED":
                    fq = vals["frequency"]
                    if fq is None or int(fq) != fq or int(fq) not in SUPPORTED_FREQUENCIES:
                        err("frequency", f"Frequency must be one of {SUPPORTED_FREQUENCIES}")
                    if vals["clean_price"] is not None and not (0 < vals["clean_price"] < 300):
                        err("clean_price", "Clean price must be in % of par, between 0 and 300")
                    cr = vals["coupon_rate"]
                    if cr is not None and not (0 <= cr <= 30):
                        err("coupon_rate", "Coupon must be given in percent, between 0 and 30")
                    elif cr is not None and 0 < cr < 0.3:
                        warn("coupon_rate", f"Coupon {cr} is read as {cr}% (percent convention); check it is not a decimal")
                    if mat is not None and mat <= as_of:
                        err("maturity", f"Maturity {mat} is not after the valuation date {as_of}")
                    if vals["nominal"] == 0:
                        err("nominal", "Nominal is zero")
                    dc = _str(rec.get("daycount")) or "ACT/ACT ICMA"
                    try:
                        dc = normalize_daycount(dc)
                    except ValueError as e:
                        err("daycount", str(e))
                    if _str(rec.get("daycount")) is None:
                        warn("daycount", "No day-count given: ACT/ACT ICMA assumed")
                    rating = _str(rec.get("rating"))
                    if not sov and rating_bucket(rating) is None:
                        warn("rating", "Missing or unrecognised rating: credit-spread risk will not be captured")
                    if status != "error":
                        pos = FixedBondPosition(
                            id=pid, nominal=vals["nominal"], clean_price=vals["clean_price"], coupon_rate=cr / 100.0,
                            frequency=int(fq), maturity=mat, currency=ccy, daycount=dc, issuer=_str(rec.get("issuer")),
                            rating=rating.upper() if rating else None, sector=_str(rec.get("sector")) or ("Sovereign" if sov else None),
                            country=_str(rec.get("country")), is_sovereign=sov, **common,
                        )
                elif itype == "FX_CASH" and status != "error":
                    pos = FxCashPosition(id=pid, currency=ccy, amount=vals["amount"], **common)
            except (TypeError, AttributeError) as e:
                err(None, f"Row could not be interpreted: {e}")
                pos = None
        if pos is not None and status != "error":
            positions.append(pos)
        preview.append({"row": r_i, "id": pid, "instrument_type": itype, "status": status})
    return positions, [i.dict() for i in issues], preview
