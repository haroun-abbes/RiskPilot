"""SYNTHETIC demonstration portfolio (fictitious issuers, consistent with the
synthetic dataset). Equity prices are the dataset closes at the valuation
date; bond clean prices are derived from yield = curve(T) + spread so that
the demo is internally consistent. Nothing here is market data."""
from __future__ import annotations

import csv
import io
from datetime import date

import numpy as np
import pandas as pd

from ..quant import bonds as bd
from ..quant.engine import curve_factors, interpolation_weights
from ..quant.instruments import EquityPosition, FixedBondPosition, FxCashPosition, Portfolio
from .serde import TEMPLATE_COLUMNS, to_template_row

AS_OF = date(2026, 10, 14)
NAME = "Demo multi-asset book (SYNTHETIC)"

EQ = [  # id, ticker, qty, type, sector, country, ccy, proxy, name
    ("EQ-NRDB", "NRDB", 40_000, "EQUITY", "Financials", "DE", "EUR", None, "Nordbank AG"),
    ("EQ-HELX", "HELX", 6_000, "EQUITY", "Health Care", "FR", "EUR", None, "Helix Pharma SE"),
    ("EQ-VLTA", "VLTA", 60_000, "EQUITY", "Utilities", "FR", "EUR", None, "Volta Energie SA"),
    ("EQ-ORBS", "ORBS", 3_000, "EQUITY", "Technology", "NL", "EUR", None, "Orbis Software NV"),
    ("EQ-CALD", "CALD", 15_000, "EQUITY", "Industrials", "IT", "EUR", None, "Caldera Industrie SpA"),
    ("EQ-MERI-S", "MERI", -8_000, "EQUITY", "Consumer Discretionary", "ES", "EUR", None, "Meridian Retail SA (short)"),
    ("EQ-ATLC", "ATLC", 1_200, "EQUITY", "Technology", "US", "USD", None, "Atlas Cloud Inc."),
    ("EQ-PNCL", "PNCL", 4_000, "EQUITY", "Health Care", "US", "USD", None, "Pinnacle Health Corp."),
    ("EQ-THMS", "THMS", 50_000, "EQUITY", "Utilities", "GB", "GBP", None, "Thames Utilities plc"),
    ("EQ-ALPN", "ALPN", 2_500, "EQUITY", "Consumer Staples", "CH", "CHF", None, "Alpine Foods AG"),
    ("EQ-KSTR", "KSTR", 5_000, "EQUITY", "Technology", "DE", "EUR", "SXEU", "Kestrel Robotics AG (proxy-mapped)"),
    ("ETF-SXEU", "SXEU", 30_000, "ETF", "Equity ETF", "EU", "EUR", None, "Synthetic Euro Equity ETF"),
    ("ETF-SXUS", "SXUS", 12_000, "ETF", "Equity ETF", "US", "USD", None, "Synthetic US Equity ETF"),
]
KSTR_PRICE = 18.40  # instrument without its own history -> proxy

BONDS = [  # id, name, issuer, nominal, cpn%, freq, maturity, dc, ccy, rating, sovereign, spread mult, sector, country
    ("BD-EUGOV-31", "EUR Sovereign 2.50% 2031", "Synthetic Republic (AAA)", 3_000_000, 2.50, 1, date(2031, 7, 4), "ACT/ACT ICMA", "EUR", "AAA", True, 0.0, "Sovereign", "EU"),
    ("BD-EUGOV-45", "EUR Sovereign 3.25% 2045", "Synthetic Republic (AAA)", 1_500_000, 3.25, 1, date(2045, 8, 15), "ACT/ACT ICMA", "EUR", "AAA", True, 0.0, "Sovereign", "EU"),
    ("BD-NRDB-29", "Nordbank 3.875% 2029", "Nordbank AG", 1_000_000, 3.875, 1, date(2029, 3, 12), "ACT/ACT ICMA", "EUR", "A-", False, 0.8, "Financials", "DE"),
    ("BD-VLTA-34", "Volta Energie 4.125% 2034", "Volta Energie SA", 1_200_000, 4.125, 1, date(2034, 11, 20), "ACT/ACT ICMA", "EUR", "BBB+", False, 1.1, "Utilities", "FR"),
    ("BD-CALD-28", "Caldera 6.50% 2028", "Caldera Industrie SpA", 800_000, 6.50, 2, date(2028, 6, 30), "30E/360", "EUR", "BB", False, 1.0, "Industrials", "IT"),
    ("BD-ATLC-32", "Atlas Cloud 5.25% 2032", "Atlas Cloud Inc.", 1_000_000, 5.25, 2, date(2032, 2, 15), "30/360", "USD", "BBB", False, 1.3, "Technology", "US"),
    ("BD-USGOV-35", "USD Sovereign 4.00% 2035", "Synthetic US Treasury", 1_000_000, 4.00, 2, date(2035, 5, 15), "ACT/ACT ICMA", "USD", "AA+", True, 0.0, "Sovereign", "US"),
]
FXC = [("CASH-EUR", "EUR", 750_000, "EUR cash"), ("CASH-USD", "USD", 400_000, "USD cash"), ("CASH-GBP", "GBP", -150_000, "GBP overdraft")]


def build(levels: pd.DataFrame) -> Portfolio:
    last = levels.loc[: pd.Timestamp(AS_OF)].iloc[-1]
    positions = []
    for pid, tk, q, typ, sec, ctry, ccy, proxy, name in EQ:
        price = KSTR_PRICE if tk == "KSTR" else round(float(last[f"EQ:{tk}"]), 2)
        positions.append(EquityPosition(pid, tk, q, price, ccy, typ, name, sec, name.split(" (")[0], ctry, proxy))
    for pid, name, issuer, nom, cpn, f, mat, dc, ccy, rating, sov, mult, sec, ctry in BONDS:
        curve = curve_factors(levels.columns, ccy)
        t = (mat - AS_OF).days / 365.25
        w = interpolation_weights([x for _, x in curve], t)
        y = float(np.dot(w, [last[c] for c, _ in curve]))
        bucket = "HY" if rating.startswith("BB") and not rating.startswith("BBB") else "IG"
        y += 0.0010 if sov else mult * float(last[f"CS:{ccy}:{bucket}"])
        cf = bd.cash_flows(bd.FixedRateBond(mat, cpn / 100, f, dc), AS_OF)
        clean = round(bd.clean_price(cf, y), 3)
        positions.append(FixedBondPosition(pid, nom, clean, cpn / 100, f, mat, ccy, dc, "BOND_FIXED", name, issuer, rating, sec, ctry, sov))
    for pid, ccy, amt, name in FXC:
        positions.append(FxCashPosition(pid, ccy, amt, "FX_CASH", name))
    return Portfolio(NAME, "EUR", AS_OF, positions)


def to_csv(pf: Portfolio) -> bytes:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=TEMPLATE_COLUMNS)
    w.writeheader()
    for p in pf.positions:
        w.writerow(to_template_row(p))
    return buf.getvalue().encode("utf-8")


def to_xlsx(pf: Portfolio) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    from .importer import FIELDS, REQUIRED

    wb = Workbook()
    ws = wb.active
    ws.title = "Positions"
    ws.append(TEMPLATE_COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1E2A44")
    for p in pf.positions:
        r = to_template_row(p)
        ws.append([r[c] if r[c] != "" else None for c in TEMPLATE_COLUMNS])
    for i, col in enumerate(TEMPLATE_COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, len(col) + 2)
    info = wb.create_sheet("Instructions")
    info.append(["RiskPilot portfolio template - SYNTHETIC demo content (fictitious issuers)"])
    info.append([])
    info.append(["Field", "Meaning"])
    for k, v in FIELDS.items():
        info.append([k, v])
    info.append([])
    info.append(["Instrument type", "Required fields"])
    for k, v in REQUIRED.items():
        info.append([k, ", ".join(v)])
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 80
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
